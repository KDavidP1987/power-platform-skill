"""Word-document furniture for an app's documentation set, plus the screenshot trim, the
inventory check and the Word finishing pass. The formatting half only: the words live in a
separate build script per project (see build_guides_example.py), so prose can be edited without
reading python-docx.

    python doc_kit.py --selftest
    python doc_kit.py trim <captures-dir> <out-dir> [--mode app|email] [--email PAT]... [--exclude PAT]...
    python doc_kit.py inventory <shots.json> [--guide NAME] [--build STAMP] [--figures FILE-OR-GLOB]...
    python doc_kit.py finish <docx-or-dir> [--no-pdf]
    python doc_kit.py render <pdf-or-dir> <out-dir> [--dpi N]

trim       crops the player chrome from app captures (never in place); files matching --email
           (default email-*) are copied as taken, files matching --exclude are skipped.
inventory  fails on a missing file, on a required (screen, role) with no shot, on a figure a build
           used that shots.json does not list (--figures, the .figures.json each save writes), and
           on a build stamp that differs from the one given (--build).
finish     opens each .docx in Word (Windows, COM), updates every field and the table of contents,
           saves, exports a PDF beside it and prints the page count. Falls back to LibreOffice
           (soffice) for the PDF; with neither it says so and exits 2.
render     turns each PDF page into a PNG (pypdfium2, else pdftoppm) so every page can be read,
           and flags pages that are blank between the running head and the footer.

Colours and the font come from the app's own theme (canvas/theme.json, the file the theme intake
writes; schema in references/documentation-set.md section 6), so a page of the guide and the
screen it describes look like the same product. Without a theme the kit uses a neutral
slate-and-blue default. Needs python-docx; trim and render need Pillow.

Pagination rules built in (references/model-driven-and-docs.md, section 4): numbered procedures
restart at 1; TOC/PAGE/NUMPAGES are fields; a page break is page-break-before on the next
paragraph (never a break paragraph that can become a blank page); table header rows repeat and
stay with the first row; rows and callouts never split; short tables stay on one page; a lead-in
paragraph stays with the table or figure it introduces; captions stay with their table or figure.
"""
import fnmatch
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Length, Pt, RGBColor

CONTENT_WIDTH = Inches(6.7)
KEEP_TOGETHER_ROWS = 8      # a table this short is never split across pages
# Figure widths by name: figure(..., width="email"). A number is inches.
WIDTHS = {"full": 6.4, "wide": 5.5, "email": 4.0, "half": 3.2, "narrow": 2.4}

# Neutral default; replaced by the app's theme when one is given.
DEFAULT_THEME = {
    "primary": "1F4E79", "primaryDark": "16385A", "accent": "0F766E", "text": "1F2933",
    "muted": "52606D", "border": "CBD2D9", "canvas": "F5F7FA", "warning": "B45309", "font": "Segoe UI",
}
# theme.json colour token -> kit colour
TOKEN_MAP = {"clrPrimary": "primary", "clrPrimaryDark": "primaryDark", "clrAccent": "accent",
             "clrText": "text", "clrTextMuted": "muted", "clrBorder": "border", "clrCanvas": "canvas",
             "clrWarning": "warning"}
# The documented schema. Anything else is reported, so a misspelt key cannot silently do nothing.
THEME_TOP_KEYS = {"status", "intake", "colours", "typography", "layout", "longText", "contrastChecked"}
THEME_COLOUR_KEYS = set(TOKEN_MAP) | {"clrSurface", "clrSuccess", "clrError", "clrInfo", "clrOnPrimary",
                                      "clrOnDark", "clrOnDarkMuted", "clrScrim"}
THEME_TYPO_KEYS = {"font", "fntBody", "scale", "weights", "note"}


def _font_name(value):
    """'Segoe UI', "Font.'Segoe UI'" or "Font.'Segoe UI' or the brand font" -> Segoe UI.
    A description ('platform default (...)') is not a font name: None."""
    value = (value or "").strip()
    if "'" in value:
        return value.split("'")[1].strip() or None
    if '"' in value:
        return value.split('"')[1].strip() or None
    if value and re.match(r"^[A-Za-z][A-Za-z0-9 \-]{0,39}$", value) and "default" not in value.lower():
        return value
    return None


def load_theme(path=None, warn=None):
    """The kit palette and font from an app theme.json; defaults otherwise.

    Reads colours.<token>.rgba [r, g, b, a] for the tokens in TOKEN_MAP and typography.font
    (typography.fntBody is read too, as the older name). Keys outside the documented schema are
    reported through warn (default: printed to stderr); keys starting with '_' are comments."""
    if warn is None:
        def warn(msg):
            sys.stderr.write("theme: %s\n" % msg)
    t = dict(DEFAULT_THEME)
    if not path or not os.path.exists(path):
        return t
    with open(path, encoding="utf-8-sig") as f:
        j = json.load(f)
    for k in j:
        if not k.startswith("_") and k not in THEME_TOP_KEYS:
            warn("unknown top-level key '%s' ignored (documented: %s)" % (k, ", ".join(sorted(THEME_TOP_KEYS))))
    colours = j.get("colours") or {}
    for k in colours:
        if not k.startswith("_") and k not in THEME_COLOUR_KEYS:
            warn("unknown colour token '%s' ignored" % k)
    for token, key in TOKEN_MAP.items():
        c = colours.get(token)
        rgba = c.get("rgba") if isinstance(c, dict) else None
        if rgba and len(rgba) >= 3 and any(rgba[:3]):
            t[key] = "%02X%02X%02X" % tuple(int(x) for x in rgba[:3])
    if "clrPrimaryDark" not in colours and t["primary"] != DEFAULT_THEME["primary"]:
        r, g, b = (int(t["primary"][i:i + 2], 16) for i in (0, 2, 4))
        t["primaryDark"] = "%02X%02X%02X" % (int(r * .72), int(g * .72), int(b * .72))
    typo = j.get("typography") or {}
    for k in typo:
        if not k.startswith("_") and k not in THEME_TYPO_KEYS:
            warn("unknown typography key '%s' ignored (the font is typography.font)" % k)
    font = _font_name(typo.get("font")) or _font_name(typo.get("fntBody"))
    if font:
        t["font"] = font
    return t


def _rgb(hex6):
    return RGBColor.from_string(hex6)


def _width(width):
    """None -> full; a preset name; a number of inches; or a python-docx Length."""
    if width is None:
        return Inches(WIDTHS["full"])
    if isinstance(width, Length):
        return width
    if isinstance(width, str):
        if width not in WIDTHS:
            raise ValueError("unknown figure width '%s' (presets: %s)" % (width, ", ".join(sorted(WIDTHS))))
        return Inches(WIDTHS[width])
    return Inches(float(width))


# --------------------------------------------------------------------- low-level XML ----
def _shade(cell, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def _cell_borders(cell, colour, size=4, sides=("top", "left", "bottom", "right")):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in sides:
        el = OxmlElement("w:" + side)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), colour)
        borders.append(el)
    tcPr.append(borders)


def _no_borders(table):
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + side)
        el.set(qn("w:val"), "none")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _row_flag(row, tag):
    """w:cantSplit or w:tblHeader on a row, once, in schema order (cantSplit before tblHeader)."""
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn(tag)) is not None:
        return
    el = OxmlElement(tag)
    if tag == "w:cantSplit" and trPr.find(qn("w:tblHeader")) is not None:
        trPr.find(qn("w:tblHeader")).addprevious(el)
    else:
        trPr.append(el)


def _keep_row_with_next(row):
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.keep_with_next = True


def _field(paragraph, code, placeholder="", dirty=True):
    """A Word field (TOC, PAGE, NUMPAGES) that the reader's Word evaluates."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    if dirty:
        begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, text, end):
        run._r.append(el)
    return run


def _rule(paragraph, colour, size=12, space=1):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(size)), ("w:space", str(space)), ("w:color", colour)):
        bottom.set(qn(k), v)
    borders.append(bottom)
    pPr.append(borders)


def _text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def _has_picture(el):
    return el.tag == qn("w:p") and el.find(".//" + qn("w:drawing")) is not None


def _is_break_paragraph(el):
    if el.tag != qn("w:p") or _text(el):
        return False
    br = el.find(".//" + qn("w:br"))
    return br is not None and br.get(qn("w:type")) == "page"


# ------------------------------------------------------------------------- the builder --
class Guide(object):
    """One document of the set: a user SOP, a manager or administrator guide, or the developer guide.

    meta is a list of (label, value) rows for the cover: audience, version, issued, app, owner.
    Put the version, the issue date and the build stamp there - a guide without them cannot be
    matched to a build. build is recorded in the .figures.json written beside the .docx, which
    `inventory --figures` checks against shots.json. With draft=True a missing screenshot becomes
    a visible "Capture needed" box instead of stopping the build.
    """

    def __init__(self, title, subtitle, blurb, meta, shots_dir, running_head, org="", theme=None,
                 build=None, draft=False):
        self.t = theme or dict(DEFAULT_THEME)
        self.doc = Document()
        self.shots = shots_dir
        self.build = build
        self.draft = draft
        self.figure_no = 0
        self.table_no = 0
        self.figures = []
        self.missing = []
        self._captions = []
        self._styles()
        self._page()
        self._running_head(running_head)
        self._cover(title, subtitle, blurb, meta, org)
        self._contents()

    def _run(self, p, text, size, colour, bold=False, italic=False):
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.color.rgb = _rgb(colour)
        r.font.name = self.t["font"]
        r.font.bold = bold
        r.italic = italic
        return r

    def _styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = self.t["font"]
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = _rgb(self.t["text"])
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.16
        for name, size, colour, before, after in (("Heading 1", 18, "primary", 22, 8),
                                                  ("Heading 2", 13.5, "primaryDark", 16, 6),
                                                  ("Heading 3", 11.5, "text", 12, 4)):
            s = st[name]
            s.font.name = self.t["font"]
            s.font.size = Pt(size)
            s.font.color.rgb = _rgb(self.t[colour])
            s.font.bold = True
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Number"):
            s = st[name]
            s.font.name = self.t["font"]
            s.font.size = Pt(10.5)
            s.font.color.rgb = _rgb(self.t["text"])
            s.paragraph_format.space_after = Pt(4)
            s.paragraph_format.line_spacing = 1.14

    def _page(self):
        for section in self.doc.sections:
            for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
                setattr(section, side, Inches(0.9))
            section.different_first_page_header_footer = True

    def _running_head(self, text):
        section = self.doc.sections[0]
        head = section.header.paragraphs[0]
        head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self._run(head, text, 8, self.t["muted"])
        _rule(head, self.t["border"], size=6, space=4)
        foot = section.footer.paragraphs[0]
        foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._run(foot, "Page ", 8, self.t["muted"])
        _field(foot, "PAGE", "1")
        self._run(foot, " of ", 8, self.t["muted"])
        _field(foot, "NUMPAGES", "1")

    def _cover(self, title, subtitle, blurb, meta, org):
        doc = self.doc
        doc.add_paragraph().paragraph_format.space_after = Pt(90)
        if org:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            self._run(p, org.upper(), 10, self.t["muted"], bold=True)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(14)
        self._run(p, title.upper(), 30, self.t["primary"], bold=True)
        _rule(p, self.t["accent"], size=18, space=6)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(16)
        self._run(p, subtitle, 17, self.t["text"])
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(48)
        self._run(p, blurb, 11.5, self.t["muted"])
        table = doc.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        _no_borders(table)
        for label, value in meta:
            row = table.add_row()
            row.cells[0].width = Inches(1.5)
            row.cells[1].width = Inches(5.2)
            self._run(row.cells[0].paragraphs[0], label.upper(), 8, self.t["muted"], bold=True)
            self._run(row.cells[1].paragraphs[0], value, 10, self.t["text"])
            for cell in row.cells:
                cell.paragraphs[0].paragraph_format.space_after = Pt(3)
        self.page_break()

    def _contents(self):
        p = self.doc.add_paragraph()
        self._run(p, "Contents", 18, self.t["primary"], bold=True)
        _rule(p, self.t["border"], size=8, space=4)
        toc = self.doc.add_paragraph()
        toc.paragraph_format.space_before = Pt(8)
        _field(toc, 'TOC \\o "1-2" \\h \\z \\u', "The contents build when this document is opened in Word.")
        self.page_break()

    # -- body ----------------------------------------------------------------------------
    def h1(self, text):
        self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def h3(self, text):
        self.doc.add_heading(text, level=3)

    def page_break(self):
        """Start the next paragraph on a new page. Written as a break marker here and turned into
        page-break-before on the next paragraph at save, so a page that is already full cannot
        leave a blank page behind it."""
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def para(self, text, bold_lead=None):
        p = self.doc.add_paragraph()
        if bold_lead:
            p.add_run(bold_lead).bold = True
        p.add_run(text)
        return p

    def _item(self, p, item):
        if isinstance(item, tuple):
            p.add_run(item[0]).bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)

    def bullets(self, items):
        for item in items:
            self._item(self.doc.add_paragraph(style="List Bullet"), item)

    def _fresh_numbering(self):
        """A new list instance restarting at 1 - otherwise every procedure continues the last one."""
        numbering = self.doc.part.numbering_part.numbering_definitions._numbering
        style_num = self.doc.styles["List Number"].element.pPr.numPr.numId.val
        abstract = numbering.num_having_numId(style_num).abstractNumId.val
        num = numbering.add_num(abstract)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return num.numId

    def steps(self, items):
        """A numbered procedure, from 1. Each item: a string, or (bold lead, remainder)."""
        num_id = self._fresh_numbering()
        for item in items:
            p = self.doc.add_paragraph(style="List Number")
            numPr = p._p.get_or_add_pPr().get_or_add_numPr()
            numPr.get_or_add_ilvl().val = 0
            numPr.get_or_add_numId().val = num_id
            self._item(p, item)

    def callout(self, heading, text, tone="info"):
        """A shaded note. Its single row never splits across pages."""
        colour = {"info": self.t["primary"], "warn": self.t["warning"], "good": self.t["accent"]}[tone]
        table = self.doc.add_table(rows=1, cols=1)
        table.autofit = False
        cell = table.rows[0].cells[0]
        cell.width = CONTENT_WIDTH
        _shade(cell, self.t["canvas"])
        _cell_borders(cell, colour, size=18, sides=("left",))
        _cell_borders(cell, self.t["border"], size=4, sides=("top", "bottom", "right"))
        _row_flag(table.rows[0], "w:cantSplit")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        self._run(p, heading, 10, colour, bold=True)
        body = cell.add_paragraph()
        body.paragraph_format.space_after = Pt(0)
        self._run(body, text, 10, self.t["text"])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return table

    def table(self, headers, rows, caption=None, widths=None):
        self.table_no += 1
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.autofit = False
        hdr = table.rows[0]
        for i, text in enumerate(headers):
            cell = hdr.cells[i]
            _shade(cell, self.t["primary"])
            _cell_borders(cell, self.t["primary"])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            self._run(p, text, 9, "FFFFFF", bold=True)
        for r, values in enumerate(rows):
            row = table.add_row()
            for i, text in enumerate(values):
                cell = row.cells[i]
                _cell_borders(cell, self.t["border"])
                if r % 2 == 1:
                    _shade(cell, self.t["canvas"])
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                self._run(p, str(text), 9.5, self.t["text"])
        if widths:
            for row in table.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
        # No row splits; the header repeats on every page and never sits alone at a page foot;
        # a short table stays on one page; the last row stays with the caption.
        for row in table.rows:
            _row_flag(row, "w:cantSplit")
        _row_flag(hdr, "w:tblHeader")
        _keep_row_with_next(hdr)
        if len(table.rows) <= KEEP_TOGETHER_ROWS:
            for row in table.rows[:-1]:
                _keep_row_with_next(row)
        p = self.doc.add_paragraph()
        self._captions.append(p._p)
        if caption:
            _keep_row_with_next(table.rows[-1])
            p.paragraph_format.space_before = Pt(3)
            self._run(p, "Table %d. %s" % (self.table_no, caption), 8.5, self.t["muted"], italic=True)
        return table

    def column_table(self, columns, live=None, caption=None):
        """The columns of one table, from the build manifest, checked against live metadata.

        columns: the manifest's list for the table; each a dict with logicalName (or schemaName or
        name), displayName, type, required, description and options ([{"label", "value"}]).
        live: the attributes read from the environment on the date the guide states - either the
        Web API's EntityDefinitions(...)/Attributes 'value' list or the same simple dicts. Choice
        values come from live when given. A column in only one of the two is marked, so the guide
        shows the drift instead of hiding it."""
        man = {}
        for c in columns:
            man[_logical(c)] = _simple_column(c)
        lv = {}
        for a in (live or []):
            s = _simple_column(a)
            if s["logical"]:
                lv[s["logical"]] = s
        rows = []
        for name in list(man) + [n for n in sorted(lv) if n not in man]:
            m, l = man.get(name), lv.get(name)
            c = l or m
            src = m or l
            notes = []
            opts = (l or {}).get("options") or (m or {}).get("options")
            if opts:
                notes.append("; ".join("%s (%s)" % (o[0], o[1]) for o in opts))
            if src.get("description"):
                notes.append(src["description"])
            if live is not None and m and not l:
                notes.append("NOT LIVE: in the manifest, not in the environment")
            if m is None:
                notes.append("LIVE ONLY: in the environment, not in the manifest")
            rows.append([src.get("display") or name, name, c.get("type") or "", "Yes" if c.get("required") else "",
                         " ".join(notes)])
        return self.table(["Column", "Logical name", "Type", "Required", "Choices and notes"], rows,
                          caption=caption, widths=[1.3, 1.4, 0.9, 0.7, 2.4])

    def lint_edges_table(self, lint, caption=None):
        """The flow write -> trigger edges from lint-flows.mjs: its --json output (a dict or the
        path to the saved file) or the --verbose text (a string or a path). Cycles, if any, are
        listed under the table as a warning - a guide must never show a loop as normal."""
        edges, cycles = _read_lint(lint)
        rows = [[e["from"], e.get("kind", ""), e.get("via", ""), e["to"]] for e in edges]
        if not rows:
            rows = [["(none)", "", "", "No flow writes a table that triggers a flow"]]
        t = self.table(["Flow", "Writes (kind)", "In action", "Can trigger"], rows, caption=caption,
                       widths=[1.8, 1.0, 1.9, 2.0])
        if cycles:
            self.callout("Loops found by lint-flows", "; ".join(cycles), "warn")
        return t

    def figure(self, filename, caption, width=None):
        """A screenshot from shots_dir with a numbered caption kept on its page. width: a preset
        name from WIDTHS ("full", "wide", "email", "half", "narrow"), inches, or a Length. A
        missing file stops the build (a guide with a hole in it is not finished), except in a
        draft, where it becomes a visible "Capture needed" box."""
        w = _width(width)
        path = os.path.join(self.shots, filename)
        self.figures.append(filename)
        if not os.path.exists(path):
            if not self.draft:
                raise SystemExit("missing screenshot: %s" % path)
            self.missing.append(filename)
            self.callout("Capture needed", "%s - %s" % (filename, caption), "warn")
            return None
        self.figure_no += 1
        holder = self.doc.add_paragraph()
        holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
        holder.paragraph_format.space_before = Pt(6)
        holder.paragraph_format.space_after = Pt(3)
        holder.paragraph_format.keep_with_next = True
        holder.add_run().add_picture(path, width=w)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        self._captions.append(p._p)
        self._run(p, "Figure %d. %s" % (self.figure_no, caption), 8.5, self.t["muted"], italic=True)
        return p

    # -- finishing -------------------------------------------------------------------------
    def _paginate(self):
        """Keep lead-ins with what they introduce, and turn break paragraphs into
        page-break-before on the next paragraph."""
        from docx.text.paragraph import Paragraph
        body = self.doc.element.body
        els = [e for e in body.iterchildren() if e.tag in (qn("w:p"), qn("w:tbl"))]
        for i, e in enumerate(els[:-1]):
            nxt = els[i + 1]
            if any(e is c for c in self._captions):
                continue            # a caption belongs to what is above it, not to what follows
            if e.tag == qn("w:p") and _text(e) and not _has_picture(e) and (nxt.tag == qn("w:tbl") or _has_picture(nxt)):
                Paragraph(e, self.doc._body).paragraph_format.keep_with_next = True
        for i, e in enumerate(els):
            if not _is_break_paragraph(e):
                continue
            j = i + 1
            while j < len(els) and els[j].tag == qn("w:p") and not _text(els[j]) and not _has_picture(els[j]):
                j += 1
            if j < len(els) and els[j].tag == qn("w:p"):
                Paragraph(els[j], self.doc._body).paragraph_format.page_break_before = True
                body.remove(e)

    def save(self, path):
        """Write the .docx and, beside it, <name>.figures.json (the build and every figure used),
        which `inventory --figures` checks against shots.json."""
        d = os.path.dirname(path)
        if d and not os.path.isdir(d):
            os.makedirs(d)
        self._paginate()
        self.doc.save(path)
        with open(os.path.splitext(path)[0] + ".figures.json", "w", encoding="utf-8") as f:
            json.dump({"document": os.path.basename(path), "build": self.build, "figures": self.figures,
                       "missing": self.missing}, f, indent=1)
        return path


def _logical(c):
    return (c.get("logicalName") or c.get("LogicalName") or c.get("schemaName") or c.get("SchemaName")
            or c.get("name") or "").lower()


def _label(v):
    if isinstance(v, dict):
        return ((v.get("UserLocalizedLabel") or {}).get("Label")
                or next((x.get("Label") for x in v.get("LocalizedLabels") or []), None))
    return v


def _simple_column(c):
    """A manifest column or a Web API attribute -> {logical, display, type, required, description, options}."""
    req = c.get("required", c.get("RequiredLevel"))
    if isinstance(req, dict):
        req = req.get("Value") in ("ApplicationRequired", "SystemRequired")
    elif isinstance(req, str):
        req = req.lower() in ("true", "yes", "required", "applicationrequired", "systemrequired")
    opts = []
    raw = c.get("options")
    if raw is None:
        raw = ((c.get("OptionSet") or c.get("GlobalOptionSet") or {}).get("Options")) or []
    for o in raw:
        if isinstance(o, dict):
            opts.append((_label(o.get("label") or o.get("Label")), o.get("value", o.get("Value"))))
        else:
            opts.append((o[0], o[1]))
    return {"logical": _logical(c), "display": _label(c.get("displayName") or c.get("DisplayName")),
            "type": c.get("type") or c.get("AttributeType") or "", "required": bool(req),
            "description": _label(c.get("description") or c.get("Description")) or "", "options": opts}


_EDGE_RE = re.compile(r"edge: '(.+?)' --(\S+) '(.*?)'--> '(.+?)'")
_CYCLE_RE = re.compile(r"^FAIL  (\S+ \[.+?\]: .+)$", re.M)    # 'FAIL  <code> [<flow>]: <message>'


def _read_lint(lint):
    if isinstance(lint, str) and os.path.exists(lint):
        with open(lint, encoding="utf-8-sig") as f:
            lint = f.read()
    if isinstance(lint, str):
        try:
            lint = json.loads(lint)
        except ValueError:
            edges = [{"from": m.group(1), "kind": m.group(2), "via": m.group(3), "to": m.group(4)}
                     for m in _EDGE_RE.finditer(lint)]
            cycles = [m.group(1) for m in _CYCLE_RE.finditer(lint)]
            return edges, cycles
    return list(lint.get("edges") or []), list(lint.get("cycles") or [])


# ------------------------------------------------------------------ screenshot trim ----
# A Power Apps capture arrives wrapped in the player's black command bar, a dark grey letterbox,
# and empty app background. Trimming it makes a screenshot look like a figure, not a monitor.
def _is_surround(px):
    r, g, b = px[:3]
    return abs(r - g) < 12 and abs(g - b) < 12 and r < 90


def _is_blank_row(im, y, w):
    px = im.load()
    row = [px[x, y] for x in range(w)]
    first = row[0]
    if min(first) < 200:
        return False
    return all(max(abs(px[i] - first[i]) for i in range(3)) <= 6 for px in row[::4])


def trim(path, out_path, foot_margin=12):
    """Crop one app capture to the application itself. Returns (before, after) sizes. Falls back
    to the whole frame when detection leaves less than 200 px. The foot margin restored below
    the last content row never reaches past the app's own bottom edge into the letterbox."""
    from PIL import Image
    im = Image.open(path).convert("RGB")
    w, h = im.size
    px = im.load()
    col = 8                                    # a column clear of any centred dialog
    top = 0
    while top < h and _is_surround(px[col, top]):
        top += 1
    bottom = h - 1
    while bottom > top and _is_surround(px[col, bottom]):
        bottom -= 1
    if bottom - top < 200:
        top, bottom = 0, h - 1
    app_bottom = bottom
    while bottom > top + 200 and _is_blank_row(im, bottom, w):
        bottom -= 1
    bottom = min(app_bottom, bottom + foot_margin)
    out = im.crop((0, top, w, bottom + 1))
    out.save(out_path)
    return (w, h), out.size


def _matches(name, patterns):
    return any(fnmatch.fnmatch(name.lower(), p.lower()) for p in patterns or ())


def trim_all(src_dir, dst_dir, pattern="*.png", mode="app", email=("email-*",), exclude=()):
    """Trim every capture into dst_dir. Never edits in place, so it can be re-run.

    mode "app": player captures are trimmed; files matching an email pattern are copied as taken
    (an element capture of a reading pane has no player chrome, and the app trim would cut it).
    mode "email": every file is copied as taken. Files matching an exclude pattern are skipped.
    Returns the number of files written."""
    if mode not in ("app", "email"):
        raise ValueError("trim mode must be 'app' or 'email', not '%s'" % mode)
    if not os.path.isdir(dst_dir):
        os.makedirs(dst_dir)
    done = 0
    for p in sorted(glob.glob(os.path.join(src_dir, pattern))):
        name = os.path.basename(p)
        dst = os.path.join(dst_dir, name)
        if _matches(name, exclude):
            print("  %-40s excluded" % name)
            continue
        if mode == "email" or _matches(name, email):
            shutil.copyfile(p, dst)
            print("  %-40s copied as taken" % name)
        else:
            before, after = trim(p, dst)
            print("  %-40s %sx%s -> %sx%s" % (name, before[0], before[1], after[0], after[1]))
        done += 1
    print("wrote %d into %s" % (done, dst_dir))
    return done


def cut_band(path, y_from, y_to, fill=(245, 247, 250)):
    """Remove a horizontal strip and close the gap - for a banner only one role can see."""
    from PIL import Image
    im = Image.open(path).convert("RGB")
    w, h = im.size
    head, tail = im.crop((0, 0, w, y_from)), im.crop((0, y_to, w, h))
    out = Image.new("RGB", (w, head.size[1] + tail.size[1]), fill)
    out.paste(head, (0, 0))
    out.paste(tail, (0, head.size[1]))
    out.save(path)
    return out.size


# ------------------------------------------------------------------ inventory check ----
def load_shots(spec_path):
    with open(spec_path, encoding="utf-8-sig") as f:
        return json.load(f)


def shots_dir(spec_path):
    """The captures folder: shotsDir resolved against the folder holding shots.json (never the
    current directory), so the build and the inventory read the same files wherever they run."""
    spec = load_shots(spec_path)
    return os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(spec_path)), spec.get("shotsDir", ".")))


def _required_pairs(need):
    """A guide's requirement: roles x screens, plus explicit pairs ([screen, role] or
    {"screen", "role"}) for screens that only some roles of that guide see."""
    pairs = [(s, r) for r in need.get("roles", []) for s in need.get("screens", [])]
    for p in need.get("pairs", []):
        pairs.append((p["screen"], p["role"]) if isinstance(p, dict) else (p[0], p[1]))
    return pairs


def _placeholder(stamp):
    return not stamp or str(stamp).strip().startswith("<")


def inventory(spec_path, guide=None, build=None, figures=None):
    """Check a shot list against what each guide needs. shots.json:
        {"shotsDir": "../out/shots-trimmed",           relative to this file
         "build": "2030-01-01 09:00Z abc1234",          the build every capture was taken on
         "shots": [{"file": "home-manager.png", "screen": "Home", "role": "Manager"}, ...],
         "required": {"user": {"roles": ["Employee"], "screens": ["Home", "My requests"]},
                      "manager": {"roles": [...], "screens": [...],
                                  "pairs": [["Approvals", "Manager"]]}, ...}}
    guide: check only that guide's requirements. build: the stamp the guides are being built for;
    shots.json must carry exactly it. figures: .figures.json files written by Guide.save; every
    figure listed there must be a shot in shots.json, and each file's build must match.
    Returns the list of problems; an empty shot list is a failure, not a pass."""
    spec = load_shots(spec_path)
    base = shots_dir(spec_path)
    shots = spec.get("shots") or []
    problems = []
    if not shots:
        problems.append("no shots listed - an empty inventory is not a pass")
    have = set()
    listed = set()
    for s in shots:
        if not os.path.exists(os.path.join(base, s["file"])):
            problems.append("missing file %s (%s as %s) in %s" % (s["file"], s.get("screen"), s.get("role"), base))
        have.add((s.get("screen"), s.get("role")))
        listed.add(s["file"])
    required = spec.get("required") or {}
    if guide is not None and guide not in required:
        problems.append("no guide '%s' in required (have: %s)" % (guide, ", ".join(sorted(required)) or "none"))
    for name, need in required.items():
        if guide is not None and name != guide:
            continue
        for screen, role in _required_pairs(need):
            if (screen, role) not in have:
                problems.append("%s guide: no shot of '%s' as %s" % (name, screen, role))
    stamp = spec.get("build")
    if build is not None:
        if _placeholder(stamp):
            problems.append("shots.json has no build stamp; the captures cannot be matched to build %s" % build)
        elif str(stamp).strip() != str(build).strip():
            problems.append("captures were taken on build '%s', not '%s' - recapture or rebuild" % (stamp, build))
    for fpath in figures or []:
        with open(fpath, encoding="utf-8-sig") as f:
            used = json.load(f)
        label = used.get("document") or os.path.basename(fpath)
        for fig in used.get("figures") or []:
            if fig not in listed:
                problems.append("%s: figure %s is not listed in shots.json" % (label, fig))
        for fig in used.get("missing") or []:
            problems.append("%s: figure %s was a capture-needed placeholder (draft build)" % (label, fig))
        if used.get("build") and not _placeholder(stamp) and str(used["build"]).strip() != str(stamp).strip():
            problems.append("%s: built for '%s' but the captures are from '%s'" % (label, used["build"], stamp))
    return problems


# ---------------------------------------------------------------- finish and render ----
FINISH_PS1 = r"""
param([string]$List, [int]$Pdf = 1)
$ErrorActionPreference = "Stop"
$files = Get-Content -Raw -Encoding UTF8 $List | ConvertFrom-Json
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {
  foreach ($f in $files) {
    $doc = $word.Documents.Open($f, $false, $false, $false)
    $doc.Fields.Update() | Out-Null
    foreach ($toc in $doc.TablesOfContents) { $toc.Update() | Out-Null }
    $doc.Fields.Update() | Out-Null
    $doc.SaveAs([ref]$f, [ref]16)
    if ($Pdf) { $outPdf = [IO.Path]::ChangeExtension($f, ".pdf"); $doc.SaveAs([ref]$outPdf, [ref]17) }
    $pages = $doc.ComputeStatistics(2)
    $doc.Close([ref]0)
    Write-Output ("PAGES`t{0}`t{1}" -f $f, $pages)
  }
} finally { $word.Quit() }
"""


def _word_available():
    if os.name != "nt":
        return False
    try:
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"Word.Application\CLSID"))
        return True
    except OSError:
        return False


def _soffice():
    for exe in ("soffice", "libreoffice"):
        p = shutil.which(exe)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.exists(p):
            return p
    return None


def office_backend():
    """'word', 'libreoffice' or None."""
    if _word_available():
        return "word"
    if _soffice():
        return "libreoffice"
    return None


NO_OFFICE = ("finish: needs Microsoft Word (Windows) or LibreOffice (soffice on PATH); neither was found. "
             "Nothing was finished. Open each .docx in Word, update fields (Ctrl+A, F9), save, and export "
             "the PDF by hand, or install LibreOffice.")


def pdf_pages(pdf):
    """Page count of a PDF (pypdfium2 when installed, otherwise a count of its page objects)."""
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(pdf)
        try:
            return len(doc)
        finally:
            doc.close()
    except ImportError:
        with open(pdf, "rb") as f:
            return len(re.findall(rb"/Type\s*/Page(?![a-zA-Z])", f.read()))


def _docx_list(target):
    if os.path.isdir(target):
        return sorted(p for p in glob.glob(os.path.join(target, "*.docx")) if not os.path.basename(p).startswith("~$"))
    return [target]


def finish(target, pdf=True, backend="auto"):
    """Finish built guides: update fields and contents, save, export PDF, count pages.
    Returns (code, results) where results is [(docx, pdf or None, pages)]; code 2 means no
    Word or LibreOffice was available and nothing was done (never a pass)."""
    files = [os.path.abspath(p) for p in _docx_list(target)]
    if not files:
        print("finish: no .docx found at %s" % target)
        return 1, []
    if backend == "auto":
        backend = office_backend()
    if backend is None:
        print(NO_OFFICE)
        return 2, []
    results = []
    if backend == "word":
        tmp = tempfile.mkdtemp(prefix="doc-kit-finish-")
        script, lst = os.path.join(tmp, "finish.ps1"), os.path.join(tmp, "files.json")
        with open(script, "w", encoding="utf-8-sig") as f:
            f.write(FINISH_PS1)
        with open(lst, "w", encoding="utf-8") as f:
            json.dump(files, f)
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script,
                            "-List", lst, "-Pdf", "1" if pdf else "0"], capture_output=True, text=True)
        shutil.rmtree(tmp, ignore_errors=True)
        if r.returncode != 0:
            print("finish: Word failed:\n" + (r.stderr or r.stdout).strip())
            return 1, []
        for line in r.stdout.splitlines():
            if line.startswith("PAGES\t"):
                _, f, n = line.split("\t")
                p = os.path.splitext(f)[0] + ".pdf"
                results.append((f, p if pdf else None, int(n)))
    else:
        exe = _soffice()
        if not pdf:
            print("finish: LibreOffice cannot update a .docx in place; only the PDF export is available.")
            return 1, []
        for f in files:
            r = subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(f), f],
                               capture_output=True, text=True)
            p = os.path.splitext(f)[0] + ".pdf"
            if r.returncode != 0 or not os.path.exists(p):
                print("finish: LibreOffice failed on %s:\n%s" % (f, (r.stderr or r.stdout).strip()))
                return 1, results
            results.append((f, p, pdf_pages(p)))
        print("finish: LibreOffice exported the PDFs; it does not refresh a Word table of contents, so "
              "check the contents page (or finish in Word).")
    for f, p, n in results:
        print("  %-48s %3d pages%s" % (os.path.basename(f), n, "  -> " + os.path.basename(p) if p else ""))
    return 0, results


def _blank_page(png, band=(0.12, 0.88), limit=0.0002):
    """True when the page has (almost) no ink between the running head and the footer."""
    from PIL import Image
    im = Image.open(png).convert("L")
    w, h = im.size
    body = im.crop((0, int(h * band[0]), w, int(h * band[1])))
    dark = sum(body.histogram()[:200])
    return dark < limit * body.size[0] * body.size[1]


def render(target, out_dir, dpi=60):
    """Each PDF page -> <out_dir>/<pdf-name>-p001.png. Returns (code, {pdf: [png, ...]}, blank)
    where blank lists (pdf, page) pages with nothing between head and foot. Code 2: no renderer."""
    pdfs = sorted(glob.glob(os.path.join(target, "*.pdf"))) if os.path.isdir(target) else [target]
    if not pdfs:
        print("render: no .pdf found at %s" % target)
        return 1, {}, []
    try:
        import pypdfium2 as pdfium
    except ImportError:
        pdfium = None
    pdftoppm = shutil.which("pdftoppm")
    if pdfium is None and not pdftoppm:
        print("render: needs pypdfium2 (pip install pypdfium2) or pdftoppm (poppler); neither was found. "
              "Nothing was rendered.")
        return 2, {}, []
    if not os.path.isdir(out_dir):
        os.makedirs(out_dir)
    pages, blank = {}, []
    for pdf in pdfs:
        stem = os.path.splitext(os.path.basename(pdf))[0]
        out = []
        if pdfium is not None:
            doc = pdfium.PdfDocument(pdf)
            try:
                for i in range(len(doc)):
                    png = os.path.join(out_dir, "%s-p%03d.png" % (stem, i + 1))
                    doc[i].render(scale=dpi / 72.0).to_pil().save(png)
                    out.append(png)
            finally:
                doc.close()
        else:
            prefix = os.path.join(out_dir, stem + "-p")
            subprocess.run([pdftoppm, "-r", str(dpi), "-png", pdf, prefix], check=True)
            for p in sorted(glob.glob(prefix + "*.png")):
                n = int(re.search(r"-p-?0*(\d+)\.png$", p).group(1))
                png = os.path.join(out_dir, "%s-p%03d.png" % (stem, n))
                if p != png:
                    os.replace(p, png)
                out.append(png)
            out.sort()
        pages[pdf] = out
        for i, png in enumerate(out):
            if _blank_page(png):
                blank.append((pdf, i + 1))
        print("  %-48s %3d page(s) -> %s" % (os.path.basename(pdf), len(out), out_dir))
    for pdf, n in blank:
        print("BLANK    %s page %d has nothing between the running head and the footer" % (os.path.basename(pdf), n))
    print("render: read every page image; look for stranded headings, figures alone on a page and tables split badly.")
    return (1 if blank else 0), pages, blank


# ------------------------------------------------------------------------ self-test ----
def selftest():
    import zipfile
    from PIL import Image, ImageDraw
    from docx.table import Table
    tmp = tempfile.mkdtemp(prefix="doc-kit-")
    fails, notes = [], []

    def check(name, ok):
        if not ok:
            fails.append(name)

    # -- trim: a fake capture: black bar, grey letterbox, a pale app with a block, then blank rows.
    raw = os.path.join(tmp, "raw")
    os.makedirs(raw)
    im = Image.new("RGB", (900, 700), (245, 247, 250))
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, 899, 39), fill=(20, 20, 20))
    d.rectangle((0, 40, 899, 79), fill=(60, 60, 60))
    d.rectangle((0, 660, 899, 699), fill=(60, 60, 60))
    d.rectangle((40, 120, 600, 400), fill=(31, 78, 121))
    im.save(os.path.join(raw, "home-user.png"))
    # Content running to 4 px above the letterbox: the foot margin must stop at the app's edge.
    im2 = Image.new("RGB", (900, 700), (245, 247, 250))
    d2 = ImageDraw.Draw(im2)
    d2.rectangle((0, 0, 899, 39), fill=(20, 20, 20))
    d2.rectangle((0, 660, 899, 699), fill=(60, 60, 60))
    d2.rectangle((40, 100, 860, 655), fill=(31, 78, 121))
    im2.save(os.path.join(raw, "list-user.png"))
    mail = Image.new("RGB", (500, 300), (255, 255, 255))
    ImageDraw.Draw(mail).rectangle((0, 0, 499, 30), fill=(40, 40, 40))   # would be "chrome" to the app trim
    mail.save(os.path.join(raw, "email-approve-manager.png"))
    Image.new("RGB", (50, 50), (0, 0, 0)).save(os.path.join(raw, "scratch-ignore.png"))
    trimmed = os.path.join(tmp, "shots")
    n = trim_all(raw, trimmed, exclude=("scratch-*",))
    w, h = Image.open(os.path.join(trimmed, "home-user.png")).size
    check("trim: chrome and blank tail removed (height %d)" % h, 300 < h < 700)
    t2 = Image.open(os.path.join(trimmed, "list-user.png")).convert("RGB")
    last = [t2.getpixel((x, t2.size[1] - 1)) for x in (8, 450)]
    check("trim: foot margin reached into the letterbox (last row %s)" % (last,), not any(_is_surround(p) for p in last))
    check("trim: clamped height should be 620, got %d" % t2.size[1], t2.size[1] == 620)
    with open(os.path.join(raw, "email-approve-manager.png"), "rb") as a, \
            open(os.path.join(trimmed, "email-approve-manager.png"), "rb") as b:
        check("trim: email capture must be copied as taken", a.read() == b.read())
    check("trim: excluded file was written", not os.path.exists(os.path.join(trimmed, "scratch-ignore.png")) and n == 3)
    try:
        trim_all(raw, trimmed, mode="crop")
        fails.append("trim: an unknown mode must be refused")
    except ValueError:
        pass

    # -- theme: documented keys, the legacy key, a description that is not a font, unknown keys.
    theme = os.path.join(tmp, "theme.json")
    with open(theme, "w") as f:
        json.dump({"_about": "x", "colours": {"clrPrimary": {"rgba": [11, 83, 148, 1]},
                                              "clrPrimaryDark": {"rgba": [8, 50, 90, 1]}, "clrPrimry": {"rgba": [1, 2, 3, 1]}},
                   "typography": {"font": "Arial", "fntBody": "Font.'Segoe UI'"}, "colors": {}}, f)
    warned = []
    t = load_theme(theme, warn=warned.append)
    check("theme: primary not read (%s)" % t["primary"], t["primary"] == "0B5394")
    check("theme: clrPrimaryDark not read (%s)" % t["primaryDark"], t["primaryDark"] == "08325A")
    check("theme: typography.font not read (%s)" % t["font"], t["font"] == "Arial")
    check("theme: expected 2 warnings (colors, clrPrimry), got %s" % warned,
          len(warned) == 2 and any("colors" in m for m in warned) and any("clrPrimry" in m for m in warned))
    with open(theme, "w") as f:
        json.dump({"typography": {"fntBody": "Font.'Segoe UI' or the brand font"}}, f)
    check("theme: legacy fntBody", load_theme(theme, warn=warned.append)["font"] == "Segoe UI")
    with open(theme, "w") as f:
        json.dump({"typography": {"font": "platform default (theme)"}}, f)
    check("theme: a description must not become the font", load_theme(theme, warn=warned.append)["font"] == "Segoe UI")
    t = load_theme(None)

    # -- build: every pagination rule, widths, helpers, the figures sidecar.
    g = Guide("Sample app", "User guide", "How to use the sample app.",
              [("Audience", "Users"), ("Version", "1.0"), ("Issued", "1 January 2030"), ("Build", "B1")],
              trimmed, "Sample app - User guide", org="Example organisation", theme=t, build="B1")
    g.h1("1. Getting started")
    g.steps(["Open the app.", ("Choose ", "New request.")])
    g.h2("1.1 A second procedure")
    g.steps(["This list must start at 1 again.", "Second step."])
    g.para("Lead-in to the callout.")
    g.callout("Note", "Callouts carry a tone.", "warn")
    g.para("Lead-in to the short table.")
    g.table(["Column", "Meaning"], [["A", "first"], ["B", "second"]], caption="A short table")
    g.para("Lead-in to the figure.")
    g.figure("home-user.png", "Home, as a user")
    g.page_break()
    g.h1("2. Reference")
    g.para("Lead-in to the long table.")
    g.table(["Key", "Value"], [["k%d" % i, "v%d" % i] for i in range(30)], caption="A long table")
    g.figure("email-approve-manager.png", "An approval email", width="email")
    g.column_table([{"logicalName": "app_status", "displayName": "Status", "type": "Choice", "required": True},
                    {"logicalName": "app_retired", "displayName": "Retired", "type": "Text"}],
                   live=[{"LogicalName": "app_status", "AttributeType": "Picklist",
                          "DisplayName": {"UserLocalizedLabel": {"Label": "Status"}},
                          "RequiredLevel": {"Value": "ApplicationRequired"},
                          "OptionSet": {"Options": [{"Value": 100000000, "Label": {"UserLocalizedLabel": {"Label": "Open"}}}]}},
                         {"LogicalName": "app_extra", "AttributeType": "String"}], caption="Columns")
    g.lint_edges_table("FAIL  Flow A\n      edge: 'Flow A' --update 'Update_row'--> 'Flow B'\n"
                       "FAIL  LOOP001 [Flow A]: Flow A -> Flow B -> Flow A\n", caption="Edges")
    g.figure("home-user.png", "A figure straight after a callout", width="half")
    try:
        g.figure("home-user.png", "x", width="huge")
        fails.append("figure: an unknown width preset must be refused")
    except ValueError:
        pass
    out = g.save(os.path.join(tmp, "out", "sample.docx"))
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
        numbering = z.read("word/numbering.xml").decode("utf-8")
    for name, ok in [("TOC field", 'TOC \\o' in xml), ("figure caption", "Figure 1. Home, as a user" in xml),
                     ("table caption", "Table 1. A short table" in xml), ("restart at 1", numbering.count("w:startOverride") >= 2),
                     ("image embedded", "graphicData" in xml),
                     ("email width preset 4.0 in", 'cx="%d"' % Inches(4.0) in xml),
                     ("no page-break paragraph left", 'w:type="page"' not in xml),
                     ("column table: live choice", "Open (100000000)" in xml),
                     ("column table: drift marked", "NOT LIVE" in xml and "LIVE ONLY" in xml),
                     ("lint edges table", "Update_row" in xml),
                     ("lint cycle callout", "LOOP001 [Flow A]" in xml)]:
        check("docx: " + name, ok)
    doc = Document(out)
    body = doc.element.body
    els = [e for e in body.iterchildren() if e.tag in (qn("w:p"), qn("w:tbl"))]

    def para_with(text):
        return next(i for i, e in enumerate(els) if e.tag == qn("w:p") and _text(e) == text)

    def keep(e):
        return e.find("./" + qn("w:pPr") + "/" + qn("w:keepNext")) is not None

    def pbb(e):
        return e.find("./" + qn("w:pPr") + "/" + qn("w:pageBreakBefore")) is not None

    for lead in ("Lead-in to the callout.", "Lead-in to the short table.", "Lead-in to the figure.", "Lead-in to the long table."):
        check("keep: '%s' is not kept with what follows" % lead, keep(els[para_with(lead)]))
    check("keep: a figure caption followed by a table must not be kept with it",
          not keep(els[para_with("Figure 2. An approval email")]))
    check("page break: Contents does not start a page", pbb(els[para_with("Contents")]))
    check("page break: chapter 2 does not start a page", pbb(els[para_with("2. Reference")]))
    check("page break: chapter 1 does not start a page", pbb(els[para_with("1. Getting started")]))
    callout = Table(els[para_with("Lead-in to the callout.") + 1], doc._body)
    check("callout: row can split", callout.rows[0]._tr.trPr.find(qn("w:cantSplit")) is not None)
    short = Table(els[para_with("Lead-in to the short table.") + 1], doc._body)
    check("caption: last row of a captioned table not kept with its caption",
          all(p.paragraph_format.keep_with_next for c in short.rows[-1].cells for p in c.paragraphs))
    long_t = Table(els[para_with("Lead-in to the long table.") + 1], doc._body)
    hdr = long_t.rows[0]
    check("header: long table header does not repeat", hdr._tr.trPr.find(qn("w:tblHeader")) is not None)
    check("header: long table header can be stranded",
          all(p.paragraph_format.keep_with_next for c in hdr.cells for p in c.paragraphs))
    check("header: long table body rows should not all be kept together",
          not long_t.rows[10].cells[0].paragraphs[0].paragraph_format.keep_with_next)
    check("header: long table last row not kept with caption",
          long_t.rows[-1].cells[0].paragraphs[0].paragraph_format.keep_with_next)
    tr = hdr._tr.trPr
    check("header: trPr order must be cantSplit then tblHeader",
          [c.tag for c in tr].index(qn("w:cantSplit")) < [c.tag for c in tr].index(qn("w:tblHeader")))
    sidecar = os.path.join(tmp, "out", "sample.figures.json")
    check("save: figures sidecar not written", os.path.exists(sidecar))
    try:
        g.figure("absent.png", "x")
        fails.append("figure: a missing screenshot must stop the build")
    except SystemExit:
        pass
    dg = Guide("Sample", "Draft", "", [], trimmed, "Sample", draft=True)
    dg.figure("absent.png", "To capture")
    dg.save(os.path.join(tmp, "out", "draft.docx"))
    check("draft: a missing figure must be recorded", dg.missing == ["absent.png"])

    # -- inventory: per guide, pairs, shotsDir relative to shots.json, build stamp, figures.
    specdir = os.path.join(tmp, "docs")
    os.makedirs(specdir)
    spec = os.path.join(specdir, "shots.json")

    def write_spec(obj):
        with open(spec, "w") as f:
            json.dump(obj, f)

    good = {"shotsDir": "../shots", "build": "B1",
            "shots": [{"file": "home-user.png", "screen": "Home", "role": "User"},
                      {"file": "email-approve-manager.png", "screen": "Email - Approve", "role": "Manager"}],
            "required": {"user": {"roles": ["User"], "screens": ["Home"]},
                         "manager": {"roles": [], "screens": [], "pairs": [["Email - Approve", "Manager"]]},
                         "admin": {"roles": ["Administrator"], "screens": ["Settings"]}}}
    write_spec(good)
    here = os.getcwd()
    os.chdir(raw)                                   # the current directory must not matter
    try:
        check("shots_dir: not resolved against shots.json (%s)" % shots_dir(spec),
              os.path.normcase(shots_dir(spec)) == os.path.normcase(os.path.normpath(trimmed)))
        check("inventory: per guide user should pass: %s" % inventory(spec, guide="user"), inventory(spec, guide="user") == [])
        check("inventory: pairs for manager should pass: %s" % inventory(spec, guide="manager"),
              inventory(spec, guide="manager") == [])
        check("inventory: admin gap not found", len(inventory(spec, guide="admin")) == 1)
        check("inventory: whole set should report only the admin gap", len(inventory(spec)) == 1)
        check("inventory: unknown guide must fail", len(inventory(spec, guide="nobody")) == 1)
        check("build: matching stamp should pass", inventory(spec, guide="user", build="B1") == [])
        check("build: different stamp must fail", len(inventory(spec, guide="user", build="B2")) == 1)
        check("figures: sidecar lists a figure shots.json lacks? (expected pass)",
              inventory(spec, guide="user", figures=[sidecar]) == [])
        del good["shots"][1]
        good["required"].pop("manager")
        write_spec(good)
        check("figures: unlisted figure not caught", any("not listed" in p for p in inventory(spec, guide="user", figures=[sidecar])))
        good["build"] = "<build stamp>"
        write_spec(good)
        check("build: placeholder stamp must fail", len(inventory(spec, guide="user", build="B1")) == 1)
        write_spec({"shotsDir": "../shots", "shots": [{"file": "gone.png", "screen": "Admin", "role": "Administrator"}],
                    "required": {"manager": {"roles": ["Manager"], "screens": ["Home"]}}})
        check("inventory: expected a missing file and a missing pair", len(inventory(spec)) == 2)
        write_spec({"shots": []})
        check("inventory: an empty shot list must fail", bool(inventory(spec)))
    finally:
        os.chdir(here)

    # -- finish and render: only with Word or LibreOffice, and a renderer; otherwise say so.
    print("selftest: simulating a machine with neither Word nor LibreOffice:")
    code, _ = finish(out, backend=None)
    check("finish: with no Office it must exit 2, not pass", code == 2)
    backend = office_backend()
    if backend is None:
        notes.append("finish/render SKIPPED: no Microsoft Word or LibreOffice on this machine (not tested, not passed)")
    else:
        code, res = finish(out)
        check("finish (%s): failed with code %s" % (backend, code), code == 0 and len(res) == 1)
        if code == 0 and res:
            _, pdf, pages = res[0]
            check("finish: no PDF written", pdf and os.path.exists(pdf))
            check("finish: page count %s does not match the PDF's %s" % (pages, pdf_pages(pdf)), pages == pdf_pages(pdf))
            check("finish: expected at least 4 pages, got %s" % pages, pages >= 4)
            with zipfile.ZipFile(out) as z:
                check("finish: contents not built", "_Toc" in z.read("word/document.xml").decode("utf-8"))
            rcode, rpages, blank = render(pdf, os.path.join(tmp, "pages"))
            if rcode == 2:
                notes.append("render SKIPPED: no pypdfium2 or pdftoppm (not tested, not passed)")
            else:
                check("render: blank pages %s" % blank, rcode == 0 and not blank)
                check("render: one PNG per page", len(rpages.get(pdf, [])) == pages)
                blank_png = os.path.join(tmp, "blank.png")
                Image.new("RGB", (510, 660), (255, 255, 255)).save(blank_png)
                check("render: a blank page is not detected", _blank_page(blank_png))
                notes.append("finish (%s) and render tested: %d pages, no blank page" % (backend, pages))
    for n in notes:
        print(n)
    if fails:
        print("selftest FAILED:\n  " + "\n  ".join(fails))
        return 1
    print("selftest ok: trim (chrome, clamp to the app edge, email copied, exclude), theme (font, legacy key, "
          "warnings), pagination (break-before, lead-ins, callout, captions, header rows), widths, column and "
          "lint-edge tables, figures sidecar, draft, inventory (per guide, pairs, shotsDir, build, figures, empty)")
    return 0


def _main(argv):
    import argparse
    if argv[:1] == ["--selftest"]:
        return selftest()
    ap = argparse.ArgumentParser(prog="doc_kit.py", description=__doc__.strip().splitlines()[0])
    sub = ap.add_subparsers(dest="cmd")
    p = sub.add_parser("trim")
    p.add_argument("src")
    p.add_argument("dst")
    p.add_argument("--mode", choices=("app", "email"), default="app")
    p.add_argument("--email", action="append", help="file pattern copied as taken (default email-*)")
    p.add_argument("--exclude", action="append", default=[])
    p = sub.add_parser("inventory")
    p.add_argument("spec")
    p.add_argument("--guide")
    p.add_argument("--build")
    p.add_argument("--figures", action="append", default=[], help="a .figures.json file or a glob")
    p = sub.add_parser("finish")
    p.add_argument("target")
    p.add_argument("--no-pdf", action="store_true")
    p = sub.add_parser("render")
    p.add_argument("target")
    p.add_argument("out")
    p.add_argument("--dpi", type=int, default=60)
    a = ap.parse_args(argv)
    if a.cmd == "trim":
        trim_all(a.src, a.dst, mode=a.mode, email=tuple(a.email or ("email-*",)), exclude=tuple(a.exclude))
        return 0
    if a.cmd == "inventory":
        figs = []
        for pat in a.figures:
            hits = sorted(glob.glob(pat))
            if not hits:
                print("MISSING  no figures file matches %s" % pat)
                return 1
            figs += hits
        probs = inventory(a.spec, guide=a.guide, build=a.build, figures=figs)
        for pr in probs:
            print("MISSING  " + pr)
        print("inventory: %s" % ("ok" if not probs else "%d problem(s)" % len(probs)))
        return 1 if probs else 0
    if a.cmd == "finish":
        return finish(a.target, pdf=not a.no_pdf)[0]
    if a.cmd == "render":
        return render(a.target, a.out, dpi=a.dpi)[0]
    ap.print_help()
    return 2


if __name__ == "__main__":
    sys.exit(_main(sys.argv[1:]))
